import io

import pytest
from openpyxl import Workbook, load_workbook

from tests.test_channels_people import channel, cms, employee

UC1 = "UC" + "a" * 22
UC2 = "UC" + "b" * 22
UC3 = "UC" + "c" * 22


@pytest.fixture
def founder(make_founder, login):
    make_founder("a@taggify.in", name="Asha")
    return login("a@taggify.in")


def send(c, text: str, preview=False, name="channels.csv"):
    url = "/api/channels/import/preview" if preview else "/api/channels/import"
    return c.post(url, files={"file": (name, text.encode(), "text/csv")})


def test_only_id_and_name_are_needed_and_nothing_optional_blocks(founder):
    giri = cms(founder, "Giri CMS")
    employee(founder, "Ananya Sharma")
    text = (
        "Channel ID,Channel name,Platform,CMS,Manager,Status\n"
        f"{UC1},Kapil Comedy 01,YouTube,giri cms,ananya sharma,Active\n"
        f"{UC2},Second,,Unknown CMS,Nobody,\n"
        f"https://www.youtube.com/channel/{UC3},Linked,,,,inactive\n"
        "1029384756,FB Page,,Giri CMS,,\n"
        f"{'UC' + 'd' * 22},,,,,\n"
        ",No id,,,,\n"
    )
    preview = send(founder, text, preview=True).json()
    assert (preview["added"], preview["errors"]) == (4, 2)
    assert founder.get("/api/channels").json()["total"] == 0  # preview saves nothing

    r = send(founder, text)
    assert r.status_code == 201, r.text
    out = r.json()
    lines = {x["line"]: x for x in out["lines"]}
    assert out["added"] == 4
    assert (lines[2]["cms"], lines[2]["manager"], lines[2]["notes"]) == (
        "Giri CMS",
        "Ananya Sharma",
        [],
    )
    # Unknown CMS and unknown manager: still added, each with a note.
    assert lines[3]["cms"] is None and lines[3]["manager"] is None and len(lines[3]["notes"]) == 2
    # A full link is read as the ID; status is kept.
    assert (lines[4]["external_id"], lines[4]["status"]) == (UC3, "inactive")
    # Digits = Facebook page; a YouTube CMS is not attached to it.
    assert lines[5]["platform"] == "facebook" and lines[5]["cms"] is None
    # The name is required, as in the New channel form; so is the ID.
    assert lines[6]["result"] == lines[7]["result"] == "error"
    assert lines[6]["notes"] == ["No channel name — line skipped."]

    rows = {c["external_id"]: c for c in founder.get("/api/channels?page_size=50").json()["items"]}
    assert rows[UC1]["cms_partner_id"] == giri and rows[UC1]["manager"]["name"] == "Ananya Sharma"
    assert rows[UC3]["status"] == "inactive"


def test_existing_and_repeated_ids_are_skipped(founder):
    channel(founder, "Already here", external_id=UC1)
    text = f"id,name\n{UC1},Changed name\n{UC2},New\n{UC2},Again\n"
    out = send(founder, text).json()
    assert (out["added"], out["exists"], out["duplicate"]) == (1, 1, 1)
    names = {c["external_id"]: c["name"] for c in founder.get("/api/channels").json()["items"]}
    assert names == {UC1: "Already here", UC2: "New"}
    # Uploading the same file again adds nothing.
    assert send(founder, text).json()["added"] == 0


def test_needs_an_id_column_and_reads_the_excel_template(founder):
    r = send(founder, "name,cms\nX,Y\n", preview=True)
    assert r.json()["missing_columns"] == ["Channel ID"]
    r = send(founder, f"Channel ID,cms\n{UC1},Y\n", preview=True)
    assert r.json()["missing_columns"] == ["Channel name"]
    assert send(founder, "name,cms\nX,Y\n").status_code == 422

    t = founder.get("/api/channels/import/template")
    assert t.status_code == 200
    ws = load_workbook(io.BytesIO(t.content)).active
    assert ws["A1"].value == "Channel ID"

    wb = Workbook()
    wb.active.append(["Channel ID", "Channel name"])
    wb.active.append([UC1, "From Excel"])
    buf = io.BytesIO()
    wb.save(buf)
    r = founder.post(
        "/api/channels/import",
        files={"file": ("c.xlsx", buf.getvalue(), "application/octet-stream")},
    )
    assert r.status_code == 201 and r.json()["added"] == 1


def test_import_is_per_company(make_founder, login, founder):
    send(founder, f"Channel ID,Channel name\n{UC1},One\n")
    make_founder("s@sociodoor.in", org_slug="sociodoor")
    s = login("s@sociodoor.in")
    assert s.get("/api/channels").json()["total"] == 0
    assert send(s, f"Channel ID,Channel name\n{UC1},One\n").json()["added"] == 1
